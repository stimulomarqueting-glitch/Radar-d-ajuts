"""Conversa amb Claude per preparar la sol·licitud d'un ajut (streaming).

- Model per defecte: claude-opus-5-5 (RADAR_MODEL), pensament adaptatiu amb resum visible.
- Cerca i lectura web del servidor d'Anthropic per consultar les bases oficials.
- Fallback del servidor (`fallbacks="default"`) si el model declina una petició.
- Historial NOMÉS D'AFEGIR: els torns es desen exactament com s'han enviat/rebut (inclosos els blocs de
  pensament) i no es reescriuen mai. Els canvis de context entren com a missatges de sistema.
"""

from __future__ import annotations

import base64
import datetime as dt
import mimetypes
from pathlib import Path
from typing import AsyncIterator

import anthropic

from . import plantilles
from .config import Config
from .db import BaseDades

MAX_TOKENS = 64000
MAX_CONTINUACIONS = 5
BETA_FALLBACK = "server-side-fallback-2026-07-01"
TIPUS_ACCEPTATS = {
    "application/pdf": "pdf",
    "text/plain": "text", "text/markdown": "text", "text/csv": "text",
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document": "docx",
    "image/png": "imatge", "image/jpeg": "imatge", "image/webp": "imatge", "image/gif": "imatge",
}


def tipus_fitxer(nom: str, mime: str | None) -> str | None:
    mime = mime or mimetypes.guess_type(nom)[0] or ""
    if nom.lower().endswith(".md"):
        mime = "text/markdown"
    return TIPUS_ACCEPTATS.get(mime)


def _text_docx(ruta: Path) -> str:
    import docx

    document = docx.Document(str(ruta))
    linies = [p.text for p in document.paragraphs if p.text.strip()]
    for taula in document.tables:
        for fila in taula.rows:
            linies.append(" | ".join(cel.text.strip() for cel in fila.cells))
    return "\n".join(linies)


def blocs_fitxers(fitxers: list[dict]) -> list[dict]:
    """Converteix els fitxers pujats en blocs de contingut per a l'API."""
    blocs = []
    for f in fitxers:
        ruta = Path(f["ruta"])
        tipus = tipus_fitxer(f["nom"], f["mime"])
        if tipus == "pdf":
            blocs.append({"type": "document", "title": f["nom"],
                          "source": {"type": "base64", "media_type": "application/pdf",
                                     "data": base64.standard_b64encode(ruta.read_bytes()).decode()}})
        elif tipus == "imatge":
            blocs.append({"type": "image", "source": {"type": "base64", "media_type": f["mime"],
                                                      "data": base64.standard_b64encode(ruta.read_bytes()).decode()}})
        elif tipus == "docx":
            blocs.append({"type": "text", "text": f"Contingut del fitxer «{f['nom']}»:\n\n{_text_docx(ruta)}"})
        elif tipus == "text":
            blocs.append({"type": "text", "text": f"Contingut del fitxer «{f['nom']}»:\n\n"
                                                  f"{ruta.read_text(encoding='utf-8', errors='replace')}"})
    return blocs


def actualitzacio_context(db: BaseDades, expedient_id: int, historial: list[dict], avui: dt.date) -> str:
    """Novetats des del darrer torn de l'assistent: data i documents editats a mà."""
    darrer = next((m for m in reversed(historial) if m["rol"] == "assistant"), None)
    parts = []
    data_darrera = darrer["creat"][:10] if darrer else None
    if data_darrera != avui.isoformat():
        parts.append(f"Avui és {avui:%d/%m/%Y}.")
    if darrer:
        for d in db.documents_actuals(expedient_id):
            if d["origen"] != "assistent" and d["creat"] > darrer["creat"]:
                parts.append(f"L'usuari ha editat el document «{d['titol']}» (versió {d['versio']}). "
                             f"Versió vigent:\n\n{d['contingut']}")
    return "\n\n".join(parts)


def missatges_api(historial: list[dict]) -> list[dict]:
    return [{"role": m["rol"], "content": m["contingut"]} for m in historial]


def _peticio(cfg: Config, sistema: str, missatges: list[dict]) -> dict:
    peticio = {
        "model": cfg.model,
        "max_tokens": MAX_TOKENS,
        "system": [{"type": "text", "text": sistema, "cache_control": {"type": "ephemeral"}}],
        "messages": missatges,
        "thinking": {"type": "adaptive", "display": "summarized"},
        "output_config": {"effort": cfg.esforc},
        "cache_control": {"type": "ephemeral"},  # memòria cau automàtica del final de la conversa
    }
    if cfg.cerca_web:
        peticio["tools"] = [
            {"type": "web_search_20260209", "name": "web_search", "max_uses": 5},
            {"type": "web_fetch_20260209", "name": "web_fetch", "max_uses": 5},
        ]
    if cfg.fallbacks:
        peticio["betas"] = [BETA_FALLBACK]
        peticio["fallbacks"] = "default"
    return peticio


def _text_visible(contingut: list[dict]) -> str:
    return "".join(b.get("text", "") for b in contingut if b.get("type") == "text")


async def torn(db: BaseDades, cfg: Config, expedient: dict, text: str, accio: str = "",
               client: anthropic.AsyncAnthropic | None = None, avui: dt.date | None = None) -> AsyncIterator[dict]:
    """Executa un torn de conversa i en va retornant els esdeveniments per a la interfície."""
    avui = avui or dt.date.today()
    client = client or anthropic.AsyncAnthropic()
    eid = expedient["id"]
    historial = db.missatges(eid)

    # 1. Missatge de l'usuari (fitxers pendents + text o instrucció de la plantilla)
    pendents = db.fitxers(eid, pendents=True)
    instruccio = plantilles.instruccio(accio) if accio else text
    if accio and text.strip():
        instruccio += f"\n\nIndicacions addicionals: {text.strip()}"
    contingut_usuari = blocs_fitxers(pendents) + [{"type": "text", "text": instruccio}]
    visible = (f"Prepara: {plantilles.PER_TIPUS[accio].titol}" + (f" — {text.strip()}" if text.strip() else "")
               if accio else text)
    if pendents:
        visible += "\n📎 " + ", ".join(f["nom"] for f in pendents)
    nous = [db.afegeix_missatge(eid, "user", contingut_usuari, visible=visible, accio=accio)]
    novetats = actualitzacio_context(db, eid, historial, avui)
    if novetats:
        nous.append(db.afegeix_missatge(eid, "system", novetats, visible=""))
    db.marca_fitxers_enviats([f["id"] for f in pendents])

    respostes: list[str] = []
    try:
        for _ in range(MAX_CONTINUACIONS):
            peticio = _peticio(cfg, expedient["context_sistema"], missatges_api(db.missatges(eid)))
            async with client.beta.messages.stream(**peticio) as stream:
                async for event in stream:
                    if event.type == "content_block_start" and event.content_block.type == "server_tool_use":
                        yield {"tipus": "eina", "nom": event.content_block.name}
                    elif event.type == "content_block_delta":
                        if event.delta.type == "thinking_delta" and event.delta.thinking:
                            yield {"tipus": "pensament", "text": event.delta.thinking}
                        elif event.delta.type == "text_delta":
                            yield {"tipus": "text", "text": event.delta.text}
                final = await stream.get_final_message()
            if final.stop_reason == "refusal":
                raise RuntimeError("L'assistent no ha pogut respondre aquesta petició. Reformula-la, si us plau.")
            contingut = final.to_dict(mode="json")["content"]
            text_resposta = _text_visible(contingut)
            db.afegeix_missatge(eid, "assistant", contingut, visible=text_resposta, accio=accio)
            nous = []  # a partir d'aquí el torn ja té resposta: no es desfà
            respostes.append(text_resposta)
            if final.stop_reason == "pause_turn":
                continue  # una eina del servidor ha fet pausa: es continua el mateix torn
            if final.stop_reason == "max_tokens":
                yield {"tipus": "avis", "text": "La resposta s'ha tallat per llargada. Demana «continua»."}
            break
    except Exception as e:
        if nous:  # cap resposta desada: es desfà el torn perquè es pugui repetir
            db.esborra_missatges(nous)
            db.marca_fitxers_pendents([f["id"] for f in pendents])
        if isinstance(e, anthropic.APIError):
            missatge = _missatge_error(e)
        elif isinstance(e, RuntimeError):
            missatge = str(e)
        else:
            missatge = f"Error inesperat ({type(e).__name__}). Torna-ho a provar."
        yield {"tipus": "error", "text": missatge}
        return

    db.toca(eid)
    if accio:
        document = "".join(respostes).strip()
        if document:
            p = plantilles.PER_TIPUS[accio]
            db.desa_document(eid, accio, p.titol, document)
            yield {"tipus": "document", "document": accio, "titol": p.titol}
    yield {"tipus": "fi"}


def _missatge_error(e: anthropic.APIError) -> str:
    if isinstance(e, anthropic.AuthenticationError):
        return "La clau ANTHROPIC_API_KEY no és vàlida."
    if isinstance(e, anthropic.RateLimitError):
        return "S'ha superat el límit de peticions de l'API. Torna-ho a provar d'aquí a un minut."
    if isinstance(e, anthropic.BadRequestError):
        return f"Petició no vàlida: {e.message}"
    if isinstance(e, anthropic.APIConnectionError):
        return "No s'ha pogut connectar amb l'API de Claude."
    if isinstance(e, anthropic.APIStatusError) and e.status_code >= 500:
        return "L'API de Claude té problemes ara mateix. Torna-ho a provar d'aquí a una estona."
    return f"Error de l'API: {e}"
