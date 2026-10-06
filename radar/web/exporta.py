"""Exporta els documents d'un expedient a Word (.docx) i a Markdown (.zip)."""

from __future__ import annotations

import io
import re
import zipfile

LINIA_TAULA = re.compile(r"^\s*\|.*\|\s*$")
SEPARADOR_TAULA = re.compile(r"^\s*\|?\s*:?-{2,}:?\s*(\|\s*:?-{2,}:?\s*)*\|?\s*$")


def _inline(paragraf, text: str) -> None:
    """**negreta**, *cursiva* i `codi` dins d'un paràgraf de Word."""
    for tros in re.split(r"(\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`)", text):
        if not tros:
            continue
        if tros.startswith("**") and tros.endswith("**"):
            paragraf.add_run(tros[2:-2]).bold = True
        elif tros.startswith("`") and tros.endswith("`"):
            paragraf.add_run(tros[1:-1]).font.name = "Consolas"
        elif tros.startswith("*") and tros.endswith("*") and len(tros) > 2:
            paragraf.add_run(tros[1:-1]).italic = True
        else:
            paragraf.add_run(tros)


def _celles(linia: str) -> list[str]:
    return [c.strip() for c in linia.strip().strip("|").split("|")]


def markdown_a_docx(document, markdown: str) -> None:
    linies = markdown.splitlines()
    i = 0
    while i < len(linies):
        linia = linies[i]
        if LINIA_TAULA.match(linia) and i + 1 < len(linies) and SEPARADOR_TAULA.match(linies[i + 1]):
            capcalera = _celles(linia)
            files = []
            i += 2
            while i < len(linies) and LINIA_TAULA.match(linies[i]):
                files.append(_celles(linies[i]))
                i += 1
            taula = document.add_table(rows=1, cols=len(capcalera))
            taula.style = "Table Grid"
            for cel, text in zip(taula.rows[0].cells, capcalera):
                cel.text = ""
                _inline(cel.paragraphs[0], text)
                for run in cel.paragraphs[0].runs:
                    run.bold = True
            for fila in files:
                cels = taula.add_row().cells
                for cel, text in zip(cels, fila + [""] * (len(capcalera) - len(fila))):
                    cel.text = ""
                    _inline(cel.paragraphs[0], text)
            continue
        m = re.match(r"^(#{1,4})\s+(.*)", linia)
        if m:
            document.add_heading(m.group(2).strip(), level=min(len(m.group(1)), 4))
        elif re.match(r"^\s*[-*+]\s+", linia):
            _inline(document.add_paragraph(style="List Bullet"), re.sub(r"^\s*[-*+]\s+", "", linia))
        elif re.match(r"^\s*\d+[.)]\s+", linia):
            _inline(document.add_paragraph(style="List Number"), re.sub(r"^\s*\d+[.)]\s+", "", linia))
        elif linia.strip().startswith(">"):
            p = document.add_paragraph(style="Quote") if "Quote" in [s.name for s in document.styles] \
                else document.add_paragraph()
            _inline(p, linia.strip().lstrip(">").strip())
        elif linia.strip() in ("---", "***"):
            document.add_paragraph("")
        elif linia.strip():
            _inline(document.add_paragraph(), linia.strip())
        i += 1


def docx(expedient: dict, convocatoria_nom: str, documents: list[dict]) -> bytes:
    import docx as python_docx

    document = python_docx.Document()
    document.add_heading(expedient["titol"], level=0)
    document.add_paragraph(f"Sol·licitud: {convocatoria_nom}")
    document.add_paragraph("Esborrany preparat amb el Radar d'ajuts de Stimulo. Revisar abans de presentar.")
    for d in documents:
        document.add_page_break()
        document.add_heading(f"{d['titol']} (v{d['versio']})", level=1)
        markdown_a_docx(document, d["contingut"])
    sortida = io.BytesIO()
    document.save(sortida)
    return sortida.getvalue()


def zip_markdown(documents: list[dict]) -> bytes:
    sortida = io.BytesIO()
    with zipfile.ZipFile(sortida, "w", zipfile.ZIP_DEFLATED) as z:
        for d in documents:
            z.writestr(f"{d['tipus']}-v{d['versio']}.md", f"# {d['titol']}\n\n{d['contingut']}\n")
    return sortida.getvalue()
