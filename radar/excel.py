"""Importador del recull d'ajuts en Excel (el primer motor de cerca del radar).

Normalitza el full 'Subvencions i ajudes':
- propaga la columna 'Nivell' (cel·les combinades / buides)
- repara el TRL que Excel ha convertit en data (p. ex. '3-7' -> 3 de juliol)
- extreu les dates dd/mm/aaaa de 'Terminis'
- detecta quines files encara no són al catàleg curat (data/convocatories.yaml)
"""

from __future__ import annotations

import datetime as dt
import re
from pathlib import Path

import yaml

COLUMNES = {
    "Nivell": "nivell",
    "Nom de la convocatòria": "nom",
    "Entitat": "entitat",
    "Fons / proveïment (origen del finançament)": "fons",
    "Terminis": "terminis",
    "Presentació convocatòria": "periodicitat",
    "Breu descripció": "descripcio",
    "Ajuda econòmica": "ajuda",
    "Tipologia de projectes": "tipologia",
    "TRL": "trl",
}
NIVELLS = {"catalunya": "catalunya", "ministeri i estat": "estat", "estat": "estat", "europa": "europa"}
RE_DATA = re.compile(r"\b(\d{1,2})/(\d{1,2})/(\d{4})\b")


def repara_trl(valor) -> str | None:
    """Excel interpreta '3-7' com a data (dia 3, mes 7). Ho desfem."""
    if valor is None:
        return None
    if isinstance(valor, (dt.datetime, dt.date)):
        return f"{valor.day}–{valor.month}"
    return str(valor).strip() or None


def trl_rang(text: str | None) -> list[int] | None:
    if not text:
        return None
    nums = [int(n) for n in re.findall(r"\b([1-9])\b", text)]
    if not nums:
        return None
    maxim = 9 if re.search(r"o m[ée]s|\+|or more|o superior", text) else max(nums)
    return [min(nums), maxim]


def dates(text: str | None) -> list[str]:
    if not text:
        return []
    sortida = []
    for d, m, a in RE_DATA.findall(str(text)):
        try:
            sortida.append(dt.date(int(a), int(m), int(d)).isoformat())
        except ValueError:
            continue
    return sortida


def llegeix(path: Path, full: str = "Subvencions i ajudes") -> list[dict]:
    import openpyxl  # dependència només per a la importació

    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb[full] if full in wb.sheetnames else wb.worksheets[0]
    capcalera = [c.value for c in ws[1]]
    files = []
    nivell_actual = None
    for idx, fila in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        registre = {}
        for nom_col, valor in zip(capcalera, fila):
            clau = COLUMNES.get((nom_col or "").strip())
            if clau:
                registre[clau] = valor
        if not registre.get("nom"):
            continue
        if registre.get("nivell"):
            nivell_actual = NIVELLS.get(str(registre["nivell"]).strip().lower(), str(registre["nivell"]).strip())
        registre["nivell"] = nivell_actual
        registre["trl"] = repara_trl(registre.get("trl"))
        registre["trl_rang"] = trl_rang(registre["trl"])
        registre["dates_detectades"] = dates(registre.get("terminis"))
        registre["excel_fila"] = idx
        files.append({k: (v.strip() if isinstance(v, str) else v) for k, v in registre.items()})
    return files


def compara(files: list[dict], cataleg_yaml: Path) -> list[dict]:
    """Files de l'Excel que encara no tenen fitxa al catàleg (per fila o per nom)."""
    with open(cataleg_yaml, encoding="utf-8") as f:
        cataleg = yaml.safe_load(f)["convocatories"]
    files_conegudes = {c.get("excel_fila") for c in cataleg if c.get("excel_fila")}
    noms = {c["nom"].strip().lower() for c in cataleg}
    return [r for r in files if r["excel_fila"] not in files_conegudes and r["nom"].lower() not in noms]


def desa(files: list[dict], desti: Path) -> None:
    desti.parent.mkdir(parents=True, exist_ok=True)
    with open(desti, "w", encoding="utf-8") as f:
        f.write("# Generat per `python -m radar importa-excel`. No editar a mà: edita l'Excel i torna a importar.\n")
        yaml.safe_dump({"files": files}, f, allow_unicode=True, sort_keys=False, width=110)
